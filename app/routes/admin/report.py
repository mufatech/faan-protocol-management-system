from app.utils.decorators import admin_required
from app.utils.auth import get_admin_user
from flask import Flask, render_template, request, redirect, url_for
from app import app, db
from app.models.forms import  Organization, Passenger, ServiceLevel, SpecialNeed
from datetime import datetime
import pandas as pd
from flask import send_file
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Image,
    Spacer,
    Paragraph
)
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter
import os
from io import BytesIO

@app.route('/passenger-report', methods=['GET', 'POST'])
@admin_required
def passenger_report():

    organizations = Organization.query.all()

    query = Passenger.query

    organization_id = request.args.get('organization_id')
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')

    # Filter by organization
    if organization_id:
        query = query.filter(Passenger.organization_id == organization_id)

    # Filter by date range
    if start_date and end_date:
        query = query.filter(
            Passenger.travel_date.between(start_date, end_date)
        )

    passengers = query.all()

    return render_template(
        'admin/passenger_report.html',
        passengers=passengers,
        organizations=organizations,
        selected_organization=organization_id,
        start_date=start_date,
        end_date=end_date
    )


# EXPORT TO EXCEL
@app.route('/export-excel')
@admin_required
def export_excel():

    query = Passenger.query

    organization_id = request.args.get('organization_id')
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')

    # ============================
    # APPLY FILTERS
    # ============================

    if organization_id:
        query = query.filter(
            Passenger.organization_id == organization_id
        )

    if start_date and end_date:
        query = query.filter(
            Passenger.travel_date.between(
                start_date,
                end_date
            )
        )

    passengers = query.all()

    # ============================
    # CURRENT USER
    # ============================

    downloader_name = admin.fullname

    download_datetime = datetime.now().strftime(
        "%d %B %Y, %I:%M %p"
    )

    # ============================
    # BUILD DATA
    # ============================

    data = []

    for p in passengers:

        organization_name = (
            p.organization.organization_name
            if p.organization else "N/A"
        )

        service_name = (
            p.service_level.service_name
            if p.service_level else "N/A"
        )

        officer_name = (
            p.registered_by.fullname
            if p.registered_by else "N/A"
        )

        special_needs = ", ".join(
            need.need_name
            for need in p.special_needs
        ) if p.special_needs else "N/A"

        data.append({
            "ID": p.id,

            "Date Registered": (
                str(p.date_created)
                if p.date_created else ""
            ),

            "Flight": p.flight or "",

            "Passenger Name": p.passenger_name or "",

            "Itinerary": p.itinerary or "",

            "Travel Date": (
                str(p.travel_date)
                if p.travel_date else ""
            ),

            "Special Need": special_needs,

            "Service Level": service_name,

            "Additional Request": (
                p.additional_request or ""
            ),

            "Organization": organization_name,

            "Officer": officer_name,
        })

    df = pd.DataFrame(data)

    # ============================
    # CREATE EXCEL FILE
    # ============================

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            index=False,
            sheet_name="Passengers",
            startrow=8
        )

        workbook = writer.book
        worksheet = writer.sheets["Passengers"]

        # ============================
        # LANDSCAPE PAGE SETUP
        # ============================

        worksheet.page_setup.orientation = "landscape"

        worksheet.page_setup.paperSize = (
            worksheet.PAPERSIZE_A4
        )

        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

        worksheet.sheet_properties.pageSetUpPr.fitToPage = True

        worksheet.page_margins.left = 0.25
        worksheet.page_margins.right = 0.25
        worksheet.page_margins.top = 0.5
        worksheet.page_margins.bottom = 0.5

        worksheet.print_options.horizontalCentered = True

        # Repeat table heading on every page
        worksheet.print_title_rows = "1:9"

        # ============================
        # LOGO
        # ============================

        logo_path = os.path.join(
            app.root_path,
            "static",
            "img",
            "banner",
            "faanlogo.png"
        )

        if os.path.exists(logo_path):

            logo = ExcelImage(logo_path)

            logo.width = 85
            logo.height = 85

            worksheet.add_image(
                logo,
                "A1"
            )

        # ============================
        # MAIN HEADER
        # ============================

        worksheet.merge_cells("B1:K1")

        worksheet["B1"] = (
            "FEDERAL AIRPORTS AUTHORITY OF NIGERIA"
        )

        worksheet["B1"].font = Font(
            size=16,
            bold=True
        )

        worksheet["B1"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        worksheet.merge_cells("B2:K2")

        worksheet["B2"] = (
            "AIRPORT PROTOCOL MANAGEMENT SYSTEM"
        )

        worksheet["B2"].font = Font(
            size=13,
            bold=True
        )

        worksheet["B2"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        worksheet.merge_cells("B3:K3")

        worksheet["B3"] = "PASSENGER REPORT"

        worksheet["B3"].font = Font(
            size=12,
            bold=True
        )

        worksheet["B3"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        # ============================
        # REPORT INFORMATION
        # ============================

        worksheet["A5"] = "Generated:"
        worksheet["B5"] = download_datetime

        worksheet["D5"] = "Period:"
        worksheet["E5"] = (
            f"{start_date or 'All'} - "
            f"{end_date or 'All'}"
        )

        worksheet["A6"] = "Total Passengers:"
        worksheet["B6"] = len(passengers)

        # ============================
        # DOWNLOADED BY
        # ============================

        worksheet["D6"] = "Downloaded By:"
        worksheet["E6"] = downloader_name

        # ============================
        # TABLE HEADER STYLE
        # ============================

        green_fill = PatternFill(
            fill_type="solid",
            start_color="06844D",
            end_color="06844D"
        )

        white_font = Font(
            bold=True,
            color="FFFFFF"
        )

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        # Row 9 = table heading

        for cell in worksheet[9]:

            cell.fill = green_fill
            cell.font = white_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

            cell.border = thin_border

        # ============================
        # STYLE DATA
        # ============================

        for row in worksheet.iter_rows(
            min_row=10,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=worksheet.max_column
        ):

            for cell in row:

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )

                cell.border = thin_border

        # ============================
        # COLUMN WIDTHS
        # ============================

        column_widths = {
            "A": 8,
            "B": 22,
            "C": 12,
            "D": 25,
            "E": 22,
            "F": 15,
            "G": 25,
            "H": 18,
            "I": 30,
            "J": 25,
            "K": 22,
        }

        for column, width in column_widths.items():

            worksheet.column_dimensions[
                column
            ].width = width

        # ============================
        # ROW HEIGHTS
        # ============================

        worksheet.row_dimensions[1].height = 28
        worksheet.row_dimensions[2].height = 23
        worksheet.row_dimensions[3].height = 23
        worksheet.row_dimensions[9].height = 35

        for row in range(
            10,
            worksheet.max_row + 1
        ):

            worksheet.row_dimensions[
                row
            ].height = 35

        # ============================
        # SIGNATURE SECTION
        # ============================

        signature_row = worksheet.max_row + 3

        worksheet.merge_cells(
            start_row=signature_row,
            start_column=1,
            end_row=signature_row,
            end_column=4
        )

        worksheet.cell(
            signature_row,
            1
        ).value = "REPORT DOWNLOADED BY"

        worksheet.cell(
            signature_row,
            1
        ).font = Font(
            bold=True,
            size=11
        )

        worksheet.cell(
            signature_row,
            1
        ).alignment = Alignment(
            horizontal="center"
        )

        worksheet.merge_cells(
            start_row=signature_row + 1,
            start_column=1,
            end_row=signature_row + 1,
            end_column=4
        )

        worksheet.cell(
            signature_row + 1,
            1
        ).value = downloader_name

        worksheet.cell(
            signature_row + 1,
            1
        ).font = Font(
            bold=True,
            size=11
        )

        worksheet.cell(
            signature_row + 1,
            1
        ).alignment = Alignment(
            horizontal="center"
        )

        worksheet.merge_cells(
            start_row=signature_row + 2,
            start_column=1,
            end_row=signature_row + 2,
            end_column=4
        )

        worksheet.cell(
            signature_row + 2,
            1
        ).value = (
            f"Downloaded: {download_datetime}"
        )

        worksheet.cell(
            signature_row + 2,
            1
        ).alignment = Alignment(
            horizontal="center"
        )

        # ============================
        # ADD USER SIGNATURE
        # ============================

        if admin.officer_signature:

            signature_path = os.path.join(
                app.root_path,
                "static",
                admin.officer_signature
            )

            if os.path.exists(signature_path):

                signature = ExcelImage(
                    signature_path
                )

                signature.width = 150
                signature.height = 60

                worksheet.add_image(
                    signature,
                    f"A{signature_row + 3}"
                )

        # ============================
        # FREEZE HEADER
        # ============================

        worksheet.freeze_panes = "A10"

        # ============================
        # PRINT AREA
        # ============================

        worksheet.print_area = (
            f"A1:K{worksheet.max_row}"
        )

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="passenger_report.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    # EXPORT TO PDF
@app.route('/export-pdf')
@admin_required
def export_pdf():

    query = Passenger.query

    organization_id = request.args.get('organization_id')
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')

    # ============================
    # APPLY FILTERS
    # ============================

    if organization_id:

        query = query.filter(
            Passenger.organization_id == organization_id
        )

    if start_date and end_date:

        query = query.filter(
            Passenger.travel_date.between(
                start_date,
                end_date
            )
        )

    passengers = query.all()

    # ============================
    # CURRENT USER
    # ============================

    downloader_name = admin.fullname

    download_datetime = datetime.now().strftime(
        "%d %B %Y, %I:%M %p"
    )

    # ============================
    # PDF
    # ============================

    output = BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),

        rightMargin=15,
        leftMargin=15,
        topMargin=20,
        bottomMargin=20
    )

    page_width, page_height = landscape(A4)

    elements = []

    # ============================
    # STYLES
    # ============================

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "FAANTitle",
        parent=styles["Title"],
        fontSize=15,
        leading=18,
        alignment=TA_CENTER,
        spaceAfter=4
    )

    system_style = ParagraphStyle(
        "SystemTitle",
        parent=styles["Heading2"],
        fontSize=11,
        leading=14,
        alignment=TA_CENTER,
        spaceAfter=3
    )

    report_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading2"],
        fontSize=10,
        leading=12,
        alignment=TA_CENTER,
        spaceAfter=5
    )

    normal_style = ParagraphStyle(
        "NormalSmall",
        parent=styles["Normal"],
        fontSize=8,
        leading=10
    )

    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontSize=6.5,
        leading=8,
        alignment=TA_CENTER
    )

    header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontSize=6.5,
        leading=8,
        alignment=TA_CENTER,
        textColor=colors.white
    )

    # ============================
    # LOGO
    # ============================

    logo_path = os.path.join(
        app.root_path,
        "static",
        "img",
        "banner",
        "faanlogo.png"
    )

    if os.path.exists(logo_path):

        logo = Image(
            logo_path,
            width=0.9 * inch,
            height=0.7 * inch
        )

        logo.hAlign = "CENTER"

        elements.append(logo)

        elements.append(
            Spacer(1, 0.05 * inch)
        )

    # ============================
    # TITLES
    # ============================

    elements.append(
        Paragraph(
            "FEDERAL AIRPORTS AUTHORITY OF NIGERIA",
            title_style
        )
    )

    elements.append(
        Paragraph(
            "AIRPORT PROTOCOL MANAGEMENT SYSTEM",
            system_style
        )
    )

    elements.append(
        Paragraph(
            "PASSENGER REPORT",
            report_style
        )
    )

    # ============================
    # REPORT INFORMATION
    # ============================

    elements.append(
        Paragraph(
            f"<b>Generated:</b> {download_datetime}",
            normal_style
        )
    )

    elements.append(
        Paragraph(
            f"<b>Period:</b> "
            f"{start_date or 'All'} - "
            f"{end_date or 'All'}",
            normal_style
        )
    )

    elements.append(
        Paragraph(
            f"<b>Total Passengers:</b> "
            f"{len(passengers)}",
            normal_style
        )
    )

    elements.append(
        Paragraph(
            f"<b>Downloaded By:</b> "
            f"{downloader_name}",
            normal_style
        )
    )

    elements.append(
        Spacer(1, 0.12 * inch)
    )

    # ============================
    # TABLE
    # ============================

    data = [

        [
            Paragraph("ID", header_style),
            Paragraph("Date Registered", header_style),
            Paragraph("Flight", header_style),
            Paragraph("Passenger Name", header_style),
            Paragraph("Itinerary", header_style),
            Paragraph("Travel Date", header_style),
            Paragraph("Special Need", header_style),
            Paragraph("Service Level", header_style),
            Paragraph("Additional Request", header_style),
            Paragraph("Organization", header_style),
            Paragraph("Officer", header_style),
        ]
    ]

    # ============================
    # PASSENGERS
    # ============================

    for p in passengers:

        organization_name = (
            p.organization.organization_name
            if p.organization else "N/A"
        )

        service_name = (
            p.service_level.service_name
            if p.service_level else "N/A"
        )

        officer_name = (
            p.registered_by.fullname
            if p.registered_by else "N/A"
        )

        special_needs = ", ".join(
            need.need_name
            for need in p.special_needs
        ) if p.special_needs else "N/A"

        data.append([

            Paragraph(
                str(p.id),
                cell_style
            ),

            Paragraph(
                str(p.date_created)
                if p.date_created else "",
                cell_style
            ),

            Paragraph(
                p.flight or "",
                cell_style
            ),

            Paragraph(
                p.passenger_name or "",
                cell_style
            ),

            Paragraph(
                p.itinerary or "",
                cell_style
            ),

            Paragraph(
                str(p.travel_date)
                if p.travel_date else "",
                cell_style
            ),

            Paragraph(
                special_needs,
                cell_style
            ),

            Paragraph(
                service_name,
                cell_style
            ),

            Paragraph(
                p.additional_request or "",
                cell_style
            ),

            Paragraph(
                organization_name,
                cell_style
            ),

            Paragraph(
                officer_name,
                cell_style
            )
        ])

    # ============================
    # COLUMN WIDTHS
    # ============================

    available_width = (
        page_width
        - pdf.leftMargin
        - pdf.rightMargin
    )

    col_widths = [

        0.35 * inch,
        0.85 * inch,
        0.55 * inch,
        1.05 * inch,
        1.00 * inch,
        0.70 * inch,
        1.05 * inch,
        0.85 * inch,
        1.35 * inch,
        1.20 * inch,
        0.90 * inch
    ]

    total_width = sum(col_widths)

    if total_width > available_width:

        scale = available_width / total_width

        col_widths = [
            width * scale
            for width in col_widths
        ]

    # ============================
    # TABLE
    # ============================

    table = Table(
        data,
        colWidths=col_widths,
        repeatRows=1,
        hAlign="CENTER"
    )

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#06844D")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.4,
                colors.grey
            ),

            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),

            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                4
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                4
            ),

            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                3
            ),

            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                3
            ),
        ])
    )

    elements.append(table)

    # ============================
    # DOWNLOADED BY SECTION
    # ============================

    elements.append(
        Spacer(1, 0.25 * inch)
    )

    elements.append(
        Paragraph(
            "<b>REPORT DOWNLOADED BY</b>",
            normal_style
        )
    )

    elements.append(
        Paragraph(
            downloader_name,
            normal_style
        )
    )

    # ============================
    # USER SIGNATURE
    # ============================

    if admin.officer_signature:

        signature_path = os.path.join(
            app.root_path,
            "static",
            admin.officer_signature
        )

        if os.path.exists(signature_path):

            signature = Image(
                signature_path,
                width=1.5 * inch,
                height=0.6 * inch
            )

            signature.hAlign = "LEFT"

            elements.append(signature)

        else:

            elements.append(
                Paragraph(
                    "Signature: ____________________",
                    normal_style
                )
            )

    else:

        elements.append(
            Paragraph(
                "Signature: ____________________",
                normal_style
            )
        )

    elements.append(
        Paragraph(
            f"Date/Time: {download_datetime}",
            normal_style
        )
    )

    # ============================
    # BUILD
    # ============================

    pdf.build(elements)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="passenger_report.pdf",
        mimetype="application/pdf"
    )